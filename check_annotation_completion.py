from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MVP_DIR = Path(r"D:\NDVI-dataset\MVP - ndvi progression")
DEFAULT_CSV = MVP_DIR / "manual_annotation_template.csv"
DEFAULT_XLSX = MVP_DIR / "manual_annotation_template.xlsx"
REPORT_MD = MVP_DIR / "annotation_completion_report.md"
SUMMARY_JSON = MVP_DIR / "annotation_completion_summary.json"

REQUIRED_FIELDS = [
    "human_disease_score_0_to_5",
    "ndvi_stress_score_0_to_5",
    "overall_progression_state",
    "recommended_action",
    "priority",
    "confidence_0_to_5",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Check manual NDVI annotation completion.")
    parser.add_argument("--input", help="CSV or XLSX annotation template path. Defaults to CSV, then XLSX.")
    return parser.parse_args()


def choose_input(input_arg: str | None) -> Path:
    if input_arg:
        return Path(input_arg)
    if DEFAULT_CSV.exists():
        return DEFAULT_CSV
    return DEFAULT_XLSX


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return [{key: clean(value) for key, value in row.items()} for row in csv.DictReader(handle)]


def read_xlsx(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Manual Annotation"] if "Manual Annotation" in workbook.sheetnames else workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [clean(value) for value in rows[0]]
    records: list[dict[str, str]] = []
    for row in rows[1:]:
        record = {headers[idx]: clean(value) for idx, value in enumerate(row) if idx < len(headers)}
        if any(record.values()):
            records.append(record)
    return records


def read_records(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".xlsx":
        return read_xlsx(path)
    return read_csv(path)


def is_missing(record: dict[str, str], field: str) -> bool:
    value = clean(record.get(field, ""))
    return value == "" or value == "unknown"


def main() -> None:
    args = parse_args()
    input_path = choose_input(args.input)
    records = read_records(input_path)

    missing_by_field = {field: [] for field in REQUIRED_FIELDS}
    ready_records: list[str] = []
    not_ready_records: list[str] = []
    completed_records = 0

    for record in records:
        record_id = clean(record.get("record_id", "")) or "<missing record_id>"
        if clean(record.get("annotation_status")) == "complete":
            completed_records += 1

        missing_fields = [field for field in REQUIRED_FIELDS if is_missing(record, field)]
        for field in missing_fields:
            missing_by_field[field].append(record_id)

        if not missing_fields and clean(record.get("annotation_status")) == "complete":
            ready_records.append(record_id)
        else:
            not_ready_records.append(record_id)

    summary = {
        "input_path": str(input_path),
        "total_records": len(records),
        "completed_records": completed_records,
        "incomplete_records": len(records) - completed_records,
        "missing_disease_score": len(missing_by_field["human_disease_score_0_to_5"]),
        "missing_ndvi_score": len(missing_by_field["ndvi_stress_score_0_to_5"]),
        "missing_overall_progression_state": len(missing_by_field["overall_progression_state"]),
        "missing_recommended_action": len(missing_by_field["recommended_action"]),
        "missing_priority": len(missing_by_field["priority"]),
        "missing_confidence": len(missing_by_field["confidence_0_to_5"]),
        "records_ready_for_level_3": len(ready_records),
        "records_not_ready_for_level_3": len(not_ready_records),
        "ready_record_ids": ready_records,
        "not_ready_record_ids": not_ready_records,
    }

    lines = [
        "# Annotation Completion Report",
        "",
        f"- Input: `{input_path}`",
        f"- Total records: **{summary['total_records']}**",
        f"- Completed records: **{summary['completed_records']}**",
        f"- Incomplete records: **{summary['incomplete_records']}**",
        f"- Missing disease score: **{summary['missing_disease_score']}**",
        f"- Missing NDVI score: **{summary['missing_ndvi_score']}**",
        f"- Missing overall progression state: **{summary['missing_overall_progression_state']}**",
        f"- Missing recommended action: **{summary['missing_recommended_action']}**",
        f"- Missing priority: **{summary['missing_priority']}**",
        f"- Missing confidence: **{summary['missing_confidence']}**",
        f"- Records ready for Level 3: **{summary['records_ready_for_level_3']}**",
        f"- Records not ready for Level 3: **{summary['records_not_ready_for_level_3']}**",
        "",
        "## Not Ready Record IDs",
        "",
    ]
    if not_ready_records:
        lines.extend(f"- `{record_id}`" for record_id in not_ready_records)
    else:
        lines.append("- none")

    REPORT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    SUMMARY_JSON.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(f"Wrote {REPORT_MD}")
    print(f"Wrote {SUMMARY_JSON}")
    print(f"Records ready for Level 3: {len(ready_records)}")


if __name__ == "__main__":
    main()
