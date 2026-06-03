from __future__ import annotations

import csv
import json
import platform
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


MVP_DIR = Path(r"D:\NDVI-dataset\MVP - ndvi progression")
DATASET_ROOT = Path(r"D:\NDVI-dataset")

CSV_PATH = MVP_DIR / "manual_annotation_template.csv"
XLSX_PATH = MVP_DIR / "manual_annotation_template.xlsx"
GUIDE_PATH = MVP_DIR / "manual_annotation_guide.md"
PATH_FIX_REPORT_PATH = MVP_DIR / "path_fix_report.md"


COLUMNS = [
    "record_id",
    "field_name",
    "field_unit_id",
    "date_observed",
    "days_since_first_field_observation",
    "raw_ndvi_path",
    "annotated_ndvi_path",
    "rgb_path",
    "has_rgb",
    "has_ndvi",
    "has_existing_annotation",
    "crop",
    "zone_id",
    "zone_description",
    "observer_name",
    "annotation_status",
    "human_disease_score_0_to_5",
    "ndvi_stress_score_0_to_5",
    "visual_progression_state",
    "ndvi_progression_state",
    "overall_progression_state",
    "visible_symptoms",
    "lesion_presence",
    "lesion_severity",
    "lesion_area_percent_estimate",
    "chlorosis_presence",
    "necrosis_presence",
    "wilting_presence",
    "weak_growth_presence",
    "bare_soil_presence",
    "missing_crop_presence",
    "weed_pressure_presence",
    "irrigation_or_drip_issue",
    "localized_or_fieldwide",
    "spread_direction",
    "confidence_0_to_5",
    "ground_truth_note",
    "farmer_note",
    "treatment_note",
    "recommended_action",
    "priority",
    "needs_ground_inspection",
    "needs_refly",
    "reviewed_by",
    "review_date",
    "notes",
]


IMAGE_PAIRS = [
    {
        "field_name": "Mallard-Avenue",
        "date_observed": "2026-05-12",
        "raw_ndvi_path": r"raw-ndvi\Strawberry_1\Mallard-Avenue-5-12-2026-orthophoto-NDVI.png",
        "annotated_ndvi_path": r"human-Annotated-ndvi\Mallard-Avenue-5-12-2026-orthophoto-NDVI-annoated-png.png",
    },
    {
        "field_name": "Mallard-Avenue",
        "date_observed": "2026-05-18",
        "raw_ndvi_path": r"raw-ndvi\Strawberry_1\Mallard-Avenue-5-18-2026-orthophoto-NDVI.tif",
        "annotated_ndvi_path": r"human-Annotated-ndvi\Mallard-Avenue-5-18-2026-orthophoto-NDVI-annotated.png",
    },
    {
        "field_name": "Mallard-Avenue",
        "date_observed": "2026-05-22",
        "raw_ndvi_path": r"raw-ndvi\Strawberry_1\Mallard-Avenue-5-22-2026-orthophoto-NDVI.png",
        "annotated_ndvi_path": r"human-Annotated-ndvi\Mallard-Avenue-5-22-2026-orthophoto-NDVI-annotated.png",
    },
    {
        "field_name": "Mallard-Avenue",
        "date_observed": "2026-05-29",
        "raw_ndvi_path": r"raw-ndvi\Strawberry_1\Mallard-Avenue-5-29-2026-orthophoto-NDVI.png",
        "annotated_ndvi_path": r"human-Annotated-ndvi\Mallard-Avenue-5-29-2026-orthophoto-NDVI-user annotated.png",
    },
    {
        "field_name": "McIntyre-Road",
        "date_observed": "2026-05-08",
        "raw_ndvi_path": r"raw-ndvi\Strawberry_1\McIntyre-Road-5-8-2026-orthophoto-NDVI.jpg",
        "annotated_ndvi_path": r"human-Annotated-ndvi\McIntyre-Road-5-8-2026-orthophoto-NDVI-annotated.png",
    },
    {
        "field_name": "McIntyre-Road",
        "date_observed": "2026-05-27",
        "raw_ndvi_path": r"raw-ndvi\Strawberry_1\McIntyre-Road-5-27-2026-orthophoto-NDVI.jpg",
        "annotated_ndvi_path": r"human-Annotated-ndvi\McIntyre-Road-5-27-2026-orthophoto-NDVI-annotated.png",
    },
]


ALLOWED_VALUES = {
    "annotation_status": ["not_started", "in_progress", "complete", "needs_review"],
    "human_disease_score_0_to_5": ["0", "1", "2", "3", "4", "5"],
    "ndvi_stress_score_0_to_5": ["0", "1", "2", "3", "4", "5"],
    "confidence_0_to_5": ["0", "1", "2", "3", "4", "5"],
    "visual_progression_state": [
        "unknown",
        "healthy_stable",
        "suspicious_early",
        "worsening_visual_only",
        "stable_no_change",
        "improving_after_treatment",
        "insufficient_data",
    ],
    "ndvi_progression_state": [
        "unknown",
        "healthy_stable",
        "worsening_ndvi_visible",
        "stable_no_change",
        "improving_after_treatment",
        "insufficient_data",
    ],
    "overall_progression_state": [
        "unknown",
        "healthy_stable",
        "suspicious_early",
        "worsening_visual_only",
        "worsening_ndvi_visible",
        "worsening_both_rgb_and_ndvi",
        "stable_no_change",
        "improving_after_treatment",
        "insufficient_data",
    ],
    "lesion_presence": ["yes", "no", "unknown"],
    "chlorosis_presence": ["yes", "no", "unknown"],
    "necrosis_presence": ["yes", "no", "unknown"],
    "wilting_presence": ["yes", "no", "unknown"],
    "weak_growth_presence": ["yes", "no", "unknown"],
    "bare_soil_presence": ["yes", "no", "unknown"],
    "missing_crop_presence": ["yes", "no", "unknown"],
    "weed_pressure_presence": ["yes", "no", "unknown"],
    "irrigation_or_drip_issue": ["yes", "no", "unknown"],
    "localized_or_fieldwide": ["localized", "fieldwide", "mixed", "unknown"],
    "priority": ["low", "medium", "high", "urgent"],
    "needs_ground_inspection": ["yes", "no", "unknown"],
    "needs_refly": ["yes", "no", "unknown"],
}

MANUAL_REQUIRED_COLUMNS = [
    "human_disease_score_0_to_5",
    "ndvi_stress_score_0_to_5",
    "overall_progression_state",
    "confidence_0_to_5",
    "recommended_action",
    "priority",
    "needs_ground_inspection",
    "needs_refly",
    "notes",
]


def parse_iso(value: str) -> date:
    y, m, d = [int(part) for part in value.split("-")]
    return date(y, m, d)


def build_rows() -> list[dict[str, str]]:
    first_dates: dict[str, date] = {}
    for pair in IMAGE_PAIRS:
        observed = parse_iso(pair["date_observed"])
        first_dates[pair["field_name"]] = min(first_dates.get(pair["field_name"], observed), observed)

    rows: list[dict[str, str]] = []
    for pair in IMAGE_PAIRS:
        observed = parse_iso(pair["date_observed"])
        field = pair["field_name"]
        record_id = f"{field}_{pair['date_observed']}".replace("-", "_")
        row = {column: "" for column in COLUMNS}
        row.update(
            {
                "record_id": record_id,
                "field_name": field,
                "field_unit_id": field,
                "date_observed": pair["date_observed"],
                "days_since_first_field_observation": str((observed - first_dates[field]).days),
                "raw_ndvi_path": pair["raw_ndvi_path"],
                "annotated_ndvi_path": pair["annotated_ndvi_path"],
                "rgb_path": "",
                "has_rgb": "no",
                "has_ndvi": "yes",
                "has_existing_annotation": "yes",
                "crop": "strawberry",
                "zone_id": "whole_field",
                "zone_description": "Whole-field Level 2 review row. Add stable repeatable zone rows later for Level 3 expansion.",
                "annotation_status": "not_started",
                "visual_progression_state": "unknown",
                "ndvi_progression_state": "unknown",
                "overall_progression_state": "unknown",
                "lesion_presence": "unknown",
                "chlorosis_presence": "unknown",
                "necrosis_presence": "unknown",
                "wilting_presence": "unknown",
                "weak_growth_presence": "unknown",
                "bare_soil_presence": "unknown",
                "missing_crop_presence": "unknown",
                "weed_pressure_presence": "unknown",
                "irrigation_or_drip_issue": "unknown",
                "localized_or_fieldwide": "unknown",
                "priority": "medium",
                "needs_ground_inspection": "unknown",
                "needs_refly": "unknown",
            }
        )
        rows.append(row)
    return rows


def write_csv(rows: list[dict[str, str]]) -> None:
    with CSV_PATH.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Manual Annotation"

    ws.append(COLUMNS)
    for row in rows:
        ws.append([row[column] for column in COLUMNS])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill("solid", fgColor="1F4E79")
    required_fill = PatternFill("solid", fgColor="FFF2CC")
    locked_fill = PatternFill("solid", fgColor="E7EEF7")
    white_font = Font(color="FFFFFF", bold=True)

    required_indexes = {COLUMNS.index(column) + 1 for column in MANUAL_REQUIRED_COLUMNS}
    for cell in ws[1]:
        cell.fill = required_fill if cell.column in required_indexes else header_fill
        cell.font = Font(color="000000", bold=True) if cell.column in required_indexes else white_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column not in required_indexes:
                cell.fill = locked_fill

    for column_idx, column_name in enumerate(COLUMNS, start=1):
        values = [column_name] + [str(row[column_name]) for row in rows]
        width = min(max(max(len(value) for value in values) + 2, 12), 55)
        ws.column_dimensions[get_column_letter(column_idx)].width = width

    for column_name, values in ALLOWED_VALUES.items():
        column_idx = COLUMNS.index(column_name) + 1
        column_letter = get_column_letter(column_idx)
        formula = '"' + ",".join(values) + '"'
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Choose a value from the dropdown list."
        validation.errorTitle = "Invalid value"
        validation.prompt = "Choose one of the allowed values."
        validation.promptTitle = column_name
        ws.add_data_validation(validation)
        validation.add(f"{column_letter}2:{column_letter}1000")

    readme = wb.create_sheet("README")
    readme_rows = [
        ["Purpose", "Convert Level 2 whole-field NDVI progression evidence into Level 3 annotated disease progression records."],
        ["Dataset root", str(DATASET_ROOT)],
        ["MVP folder", str(MVP_DIR)],
        ["Rule", "Track visual disease and NDVI stress separately."],
        ["Required fields", ", ".join(MANUAL_REQUIRED_COLUMNS)],
        ["Level 3 minimum", "Complete all 6 current field/date rows, then add repeatable zone rows for Mallard-Avenue and McIntyre-Road."],
    ]
    for row in readme_rows:
        readme.append(row)
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 100
    for row in readme.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    scoring = wb.create_sheet("Scoring Guide")
    scoring_rows = [
        ["Score", "Disease score 0-5", "NDVI stress score 0-5", "Confidence 0-5"],
        [0, "No visible disease signal", "No NDVI stress signal", "No confidence or unusable evidence"],
        [1, "Trace or ambiguous symptoms", "Trace or ambiguous NDVI stress", "Very low confidence"],
        [2, "Mild localized symptoms", "Mild localized NDVI stress", "Low confidence"],
        [3, "Moderate clear symptoms", "Moderate clear NDVI stress", "Moderate confidence"],
        [4, "Severe broad symptoms", "Severe broad NDVI stress", "High confidence"],
        [5, "Extreme symptoms or likely crop-loss area", "Extreme NDVI stress signal", "Very high confidence"],
        ["Rule", "If symptoms visually worsen but NDVI does not, set overall_progression_state to worsening_visual_only.", "", ""],
    ]
    for row in scoring_rows:
        scoring.append(row)
    for cell in scoring[1]:
        cell.fill = header_fill
        cell.font = white_font
    for column_idx in range(1, 5):
        scoring.column_dimensions[get_column_letter(column_idx)].width = 34
    for row in scoring.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(XLSX_PATH)


def write_guide() -> None:
    GUIDE_PATH.write_text(
        """# Manual Annotation Guide

## 1. Purpose

Convert Level 2 whole-field NDVI progression into Level 3 annotated disease progression.

## 2. Why This Is Needed

The audit found repeated NDVI dates and matching human annotated overlays, but the annotations are baked into PNG files and are not machine-readable. This template captures reviewer labels, scores, progression states, and notes in a structured format that can be audited and expanded into repeatable zone labels.

## 3. Scoring

### Disease score 0-5

- 0: no visible disease signal
- 1: trace or ambiguous visible symptoms
- 2: mild localized symptoms
- 3: moderate clear symptoms
- 4: severe broad symptoms
- 5: extreme symptoms or likely crop-loss area

### NDVI stress score 0-5

- 0: no NDVI stress signal
- 1: trace or ambiguous NDVI stress
- 2: mild localized NDVI stress
- 3: moderate clear NDVI stress
- 4: severe broad NDVI stress
- 5: extreme NDVI stress signal

### Confidence 0-5

- 0: unusable or no confidence
- 1: very low confidence
- 2: low confidence
- 3: moderate confidence
- 4: high confidence
- 5: very high confidence

## 4. Important Chimera Rule

Track visual disease and NDVI stress separately.

If symptoms visually worsen but NDVI does not, mark:

`overall_progression_state = worsening_visual_only`

## 5. Minimum Work To Reach Level 3

- Complete all 6 current field/date rows.
- Then add at least 15 repeatable zones for Mallard-Avenue across 3 dates.
- Add at least 10 repeatable zones for McIntyre-Road across 2 dates.
- Assign stable `zone_id` values.
- Add 0-5 scores and notes.

## 6. Current Six Rows

- Mallard-Avenue 2026-05-12
- Mallard-Avenue 2026-05-18
- Mallard-Avenue 2026-05-22
- Mallard-Avenue 2026-05-29
- McIntyre-Road 2026-05-08
- McIntyre-Road 2026-05-27
""",
        encoding="utf-8",
    )


def write_path_fix_report() -> None:
    report = {
        "path_fix_status": "implemented",
        "dataset_resolution_priority": [
            "--dataset",
            "NDVI_DATASET_PATH",
            r"D:\NDVI-dataset",
            "/mnt/d/NDVI-dataset",
        ],
        "outdir_resolution_priority": ["--outdir", "<dataset>/MVP - ndvi progression"],
        "windows_powershell_status": "confirmed working with bundled Python and --dataset D:\\NDVI-dataset; rasterio was unavailable, so georeference probing runs in reduced mode",
        "wsl_status": "confirmed working with python3 and --dataset /mnt/d/NDVI-dataset",
        "validation_output_windows": r"D:\NDVI-dataset\MVP - ndvi progression\path_fix_validation_output_windows",
        "validation_output_wsl": r"D:\NDVI-dataset\MVP - ndvi progression\path_fix_validation_output",
        "platform_for_artifact_generation": platform.platform(),
    }
    PATH_FIX_REPORT_PATH.write_text(
        "# Path Fix Report\n\n"
        "## Status\n\n"
        "- Path handling implemented in `audit_progression_dataset.py`.\n"
        "- Dataset path priority: `--dataset`, `NDVI_DATASET_PATH`, `D:\\NDVI-dataset`, `/mnt/d/NDVI-dataset`.\n"
        "- Output folder priority: `--outdir`, then `<dataset>\\MVP - ndvi progression`.\n"
        "- Startup diagnostics now print resolved paths, platform, and required folder existence checks.\n\n"
        "## Validation\n\n"
        "- Windows PowerShell validation succeeded with `--dataset D:\\NDVI-dataset`.\n"
        "- The available Windows Python runtime does not include `rasterio`, so raster CRS/georeference probing runs in reduced mode there.\n"
        "- WSL remains the full georeference-capable validation path because `rasterio` is available there.\n"
        "- WSL validation succeeded with:\n\n"
        "```powershell\n"
        "wsl.exe -d Ubuntu-22.04 -- bash -lc \"cd '/mnt/d/NDVI-dataset/MVP - ndvi progression' && python3 audit_progression_dataset.py --dataset '/mnt/d/NDVI-dataset' --outdir '/mnt/d/NDVI-dataset/MVP - ndvi progression/path_fix_validation_output'\"\n"
        "```\n\n"
        "## Machine Summary\n\n"
        "```json\n"
        f"{json.dumps(report, indent=2)}\n"
        "```\n",
        encoding="utf-8",
    )


def main() -> None:
    rows = build_rows()
    write_csv(rows)
    write_xlsx(rows)
    write_guide()
    write_path_fix_report()
    print(f"Wrote {CSV_PATH}")
    print(f"Wrote {XLSX_PATH}")
    print(f"Wrote {GUIDE_PATH}")
    print(f"Wrote {PATH_FIX_REPORT_PATH}")
    print(f"Rows created: {len(rows)}")


if __name__ == "__main__":
    main()
